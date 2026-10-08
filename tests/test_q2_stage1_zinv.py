"""``compare_pptx_files_zinv`` on decks built with python-pptx (registration G0 item 8).

The real ``compare_pptx_files`` lives in the pinned OSWorld tree; here an order-sensitive
stand-in (shape by shape in document order, as the original zips them) shows what the
realignment changes and what it leaves alone. The validation on the stored confirm mutants
and golds runs with the pinned comparator on the host (``zinv_validate``).
"""

from __future__ import annotations

import types
from pathlib import Path

import pytest

pptx = pytest.importorskip("pptx")
from pptx.util import Emu  # noqa: E402

from harness.q2_stage1 import zinv  # noqa: E402

CM = 360_000


def deck(path: Path, shapes: list[tuple[str, int, int, int, int]], rot: dict[str, float] | None
         = None) -> Path:  # fmt: skip
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    for text, left, top, width, height in shapes:
        box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
        box.text_frame.text = text
        if rot and text in rot:
            box.rotation = rot[text]
    prs.save(path)
    return path


def order_sensitive(file1: str, file2: str, **_: object) -> int:
    """Stand-in for compare_pptx_files: shapes zipped in document order."""
    a, b = pptx.Presentation(file1), pptx.Presentation(file2)
    for s1, s2 in zip(a.slides, b.slides, strict=True):
        if len(s1.shapes) != len(s2.shapes):
            return 0
        for x, y in zip(s1.shapes, s2.shapes, strict=True):
            if (x.text_frame.text, x.left, x.top, x.width, x.height) != (
                y.text_frame.text, y.left, y.top, y.width, y.height):  # fmt: skip
                return 0
    return 1


A = ("A", 1 * CM, 1 * CM, 4 * CM, 2 * CM)
B = ("B", 12 * CM, 1 * CM, 4 * CM, 2 * CM)  # far from A
C = ("C", 2 * CM, 2 * CM, 4 * CM, 2 * CM)  # overlaps A
NEAR = ("N", 5 * CM + 50_000, 1 * CM, 4 * CM, 2 * CM)  # 0.14 cm from A: inside the clearance


def test_swap_of_apart_shapes_is_ignored(tmp_path):
    ref = deck(tmp_path / "ref.pptx", [A, B, C])
    agent = deck(tmp_path / "agent.pptx", [B, A, C])
    assert order_sensitive(str(agent), str(ref)) == 0
    assert zinv.compare_pptx_files_zinv(str(agent), str(ref), _original=order_sensitive) == 1


def test_swap_across_an_overlapping_shape_still_fails(tmp_path):
    ref = deck(tmp_path / "ref.pptx", [A, B, C])
    agent = deck(tmp_path / "agent.pptx", [C, B, A])  # A moves past C, which it overlaps
    assert zinv.compare_pptx_files_zinv(str(agent), str(ref), _original=order_sensitive) == 0


def test_shapes_within_two_millimetres_are_not_apart(tmp_path):
    ref = deck(tmp_path / "ref.pptx", [A, NEAR])
    agent = deck(tmp_path / "agent.pptx", [NEAR, A])
    assert zinv.compare_pptx_files_zinv(str(agent), str(ref), _original=order_sensitive) == 0


def test_rotation_widens_the_box(tmp_path):
    long_a = ("A", 1 * CM, 5 * CM, 12 * CM, 1 * CM)
    below = ("B", 6 * CM, 8 * CM, 2 * CM, 1 * CM)  # 2 cm below A's frame, inside its rotation
    ref = deck(tmp_path / "ref.pptx", [long_a, below], rot={"A": 45.0})
    agent = deck(tmp_path / "agent.pptx", [below, long_a], rot={"A": 45.0})
    keys = [zinv.shape_key(s) for s in pptx.Presentation(ref).slides[0].shapes]
    assert not zinv.apart(*keys)
    assert zinv.compare_pptx_files_zinv(str(agent), str(ref), _original=order_sensitive) == 0
    ref2 = deck(tmp_path / "ref2.pptx", [long_a, below])
    agent2 = deck(tmp_path / "agent2.pptx", [below, long_a])
    assert zinv.compare_pptx_files_zinv(str(agent2), str(ref2), _original=order_sensitive) == 1


def test_unchanged_order_runs_the_original_on_the_original_file(tmp_path):
    ref = deck(tmp_path / "ref.pptx", [A, B])
    agent = deck(tmp_path / "agent.pptx", [A, B])
    seen = []

    def spy(file1: str, file2: str, **options: object) -> str:
        seen.append((file1, file2, options))
        return "original-verdict"

    out = zinv.compare_pptx_files_zinv(str(agent), str(ref), _original=spy, examine_shape=False)
    assert out == "original-verdict"
    assert seen == [(str(agent), str(ref), {"examine_shape": False})]


def test_a_changed_shape_is_not_realigned(tmp_path):
    ref = deck(tmp_path / "ref.pptx", [A, B])
    moved_b = ("B", 12 * CM, 3 * CM, 4 * CM, 2 * CM)
    agent = deck(tmp_path / "agent.pptx", [moved_b, A])
    assert zinv.aligned_deck(str(agent), str(ref)) is None


def test_placeholder_without_its_own_frame_never_moves(tmp_path):
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # title placeholder inherits its frame
    title = slide.shapes.title
    assert zinv.shape_box(title._element) is None
    key = zinv.shape_key(title)
    other = zinv.ShapeKey(signature=("x",), box=(0, 0, 1, 1))
    assert not zinv.apart(key, other)


def test_align_order_keeps_identical_shapes_in_order():
    k = zinv.ShapeKey(signature=("same",), box=(0, 0, 10, 10))
    far = zinv.ShapeKey(signature=("far",), box=(10**7, 0, 10**7 + 10, 10))
    assert zinv.align_order([k, k, far], [k, k, far]) == [0, 1, 2]
    assert zinv.align_order([far, k], [k, far]) == [1, 0]
    assert zinv.align_order([k], [k, far]) is None


def test_none_paths_pass_through():
    seen = []
    zinv.compare_pptx_files_zinv(None, "ref", _original=lambda *a, **k: seen.append(a))
    assert seen == [(None, "ref")]


def test_install_replaces_only_the_metrics_namespace_entry():
    def original(*a, **k):
        return "orig"

    slides = types.SimpleNamespace(compare_pptx_files=original)

    def tolerant(f1, f2, **options):
        return slides.compare_pptx_files(f1, f2, **options)

    metrics = types.SimpleNamespace(
        compare_pptx_files=original, compare_pptx_files_tolerant=tolerant
    )
    assert zinv.install(metrics) is original
    assert metrics.compare_pptx_files is zinv.compare_pptx_files_zinv
    assert metrics.compare_pptx_files_tolerant("a", "b") == "orig"  # not corrected


def test_clearance_matches_the_mutation_operator():
    from harness.q2_mutation.operators import pptx as operator

    assert zinv.CLEARANCE_EMU == operator.CLEARANCE_EMU
