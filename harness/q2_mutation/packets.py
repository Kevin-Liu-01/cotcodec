"""Candidate artifacts for blind rater packets: structure, difference, render.

A rater sees the instruction, the initial files and the candidate (decision D9,
``raters.py``). This module turns files into what the rater reads:

* ``structure_lines``: a plain-text listing of a document's content and
  formatting (cells with formulas and cached values and number formats;
  paragraphs with style and run formatting; slides with shapes, positions and
  text; table cells). It is a display aid, not a checker: no rule, no
  comparison with gold;
* ``diff_lines``: a unified difference of the candidate's listing against the
  initial file's listing (never against gold);
* ``render_command``: the LibreOffice headless PDF conversion and
  ``pdftoppm`` page images run in the LO-VM image (the VM's own renderer).

Listings are capped so a packet stays readable.
"""

from __future__ import annotations

import argparse
import difflib
import hashlib
import json
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

MAX_LINES = 400
TEXT_SUFFIXES = frozenset({".txt", ".csv", ".json", ".py", ".md", ".html", ".css", ".js", ".xml"})


def _cap(lines: list[str], limit: int = MAX_LINES) -> list[str]:
    if len(lines) <= limit:
        return lines
    return lines[:limit] + [f"... {len(lines) - limit} more lines"]


def xlsx_lines(path: Path) -> list[str]:
    import openpyxl

    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    lines: list[str] = []
    for sheet in formulas.worksheets:
        cached = values[sheet.title]
        merged = sorted(map(str, sheet.merged_cells.ranges))
        lines.append(f"[sheet {sheet.title}] freeze={sheet.freeze_panes} merged={merged}")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                shown = cached[cell.coordinate].value
                formula = (
                    cell.value
                    if isinstance(cell.value, str) and cell.value.startswith("=")
                    else None
                )
                text = f"{sheet.title}!{cell.coordinate} = {shown!r}"
                if formula:
                    text += f" formula {formula}"
                if cell.number_format and cell.number_format != "General":
                    text += f" fmt {cell.number_format}"
                font = cell.font
                flags = [n for n, on in (("bold", font.b), ("italic", font.i)) if on]
                if flags:
                    text += " " + ",".join(flags)
                lines.append(text)
    return lines


def docx_lines(path: Path) -> list[str]:
    import docx

    document = docx.Document(str(path))
    lines: list[str] = []
    for index, paragraph in enumerate(document.paragraphs):
        runs = []
        for run in paragraph.runs:
            flags = [
                n for n, on in (("b", run.bold), ("i", run.italic), ("u", run.underline)) if on
            ]
            size = run.font.size.pt if run.font.size else None
            marks = ("/" + ",".join(flags) if flags else "") + (f"/{size}pt" if size else "")
            runs.append(f"{run.text!r}{marks}")
        align = paragraph.alignment
        lines.append(f"P{index} [{paragraph.style.name}] align={align} " + " ".join(runs))
    for t_index, table in enumerate(document.tables):
        for r_index, row in enumerate(table.rows):
            for c_index, cell in enumerate(row.cells):
                lines.append(f"T{t_index}R{r_index}C{c_index} {cell.text!r}")
    return lines


def pptx_lines(path: Path) -> list[str]:
    import pptx

    presentation = pptx.Presentation(str(path))
    lines: list[str] = []
    for s_index, slide in enumerate(presentation.slides, 1):
        lines.append(f"[slide {s_index}] layout={slide.slide_layout.name}")
        for shape in slide.shapes:
            box = f"@({shape.left},{shape.top},{shape.width},{shape.height})"
            text = shape.text_frame.text if getattr(shape, "has_text_frame", False) else ""
            lines.append(f"S{s_index} {shape.shape_type} {shape.name!r} {box} {text!r}")
            if getattr(shape, "has_table", False):
                for r_index, row in enumerate(shape.table.rows):
                    for c_index, cell in enumerate(row.cells):
                        lines.append(f"S{s_index} table R{r_index}C{c_index} {cell.text!r}")
    return lines


def structure_lines(path: Path) -> list[str]:
    """Plain-text listing of one file; unknown types get size and hash."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            return _cap(xlsx_lines(path))
        if suffix == ".docx":
            return _cap(docx_lines(path))
        if suffix == ".pptx":
            return _cap(pptx_lines(path))
        if suffix in TEXT_SUFFIXES:
            return _cap(path.read_text(encoding="utf-8", errors="replace").splitlines())
    except Exception as exc:  # noqa: BLE001 - an unreadable file is shown as such
        return [f"unreadable {suffix} file: {type(exc).__name__}: {str(exc)[:200]}"]
    data = path.read_bytes()
    return [
        f"binary {suffix or 'file'}: {len(data)} bytes, sha256 {hashlib.sha256(data).hexdigest()}"
    ]


def diff_lines(initial: Sequence[str], candidate: Sequence[str]) -> list[str]:
    """Unified difference of the candidate listing against the initial listing."""
    return _cap(
        list(
            difflib.unified_diff(
                list(initial), list(candidate), "initial", "candidate", lineterm="", n=1
            )
        )
    )


def artifacts(
    initial: Mapping[str, str], candidate: Mapping[str, str | None]
) -> dict[str, dict[str, Any]]:
    """Per VM path: candidate structure and its difference against the initial file."""
    out: dict[str, dict[str, Any]] = {}
    for vm_path in sorted(candidate):
        local = candidate[vm_path]
        if local is None:
            out[vm_path] = {"structure": ["file absent in the end state"], "diff_vs_initial": []}
            continue
        lines = structure_lines(Path(local))
        before = structure_lines(Path(initial[vm_path])) if vm_path in initial else []
        out[vm_path] = {
            "structure": lines,
            "diff_vs_initial": diff_lines(before, lines) if vm_path in initial else ["new file"],
        }
    return out


def render_command(document: str, outdir: str, profile_home: str) -> list[list[str]]:
    """Commands (LO-VM image) that render a document to PNG pages at 50 dpi."""
    stem = Path(document).stem
    return [
        [
            "env",
            f"HOME={profile_home}",
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            outdir,
            document,
        ],
        ["pdftoppm", "-r", "50", "-l", "10", "-png", f"{outdir}/{stem}.pdf", f"{outdir}/{stem}"],
    ]


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--items", type=Path, required=True, help="JSONL: item_id, initial, candidate"
    )
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    rows = []
    for line in args.items.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        item = json.loads(line)
        rows.append(
            {"item_id": item["item_id"], "candidate": artifacts(item["initial"], item["candidate"])}
        )
    args.out.write_text(
        "".join(json.dumps(r, sort_keys=True) + "\n" for r in rows), encoding="utf-8"
    )
    print(json.dumps({"items": len(rows)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
