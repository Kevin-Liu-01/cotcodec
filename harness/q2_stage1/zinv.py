"""``compare_pptx_files_zinv``: OSWorld's ``compare_pptx_files``, invariant to the stacking
order of shapes that cannot paint the same pixel (q2-stage1-rescoped-v1, G0 item 8, section 8).

The checker-mutation study (``q2-evaluator-mutation-v1``, D35) found that
``compare_pptx_files`` zips the two decks' top-level shapes in document order, so swapping the
z-order of two shapes that do not overlap, which changes nothing a viewer sees, fails the
checker (``pptx.eq.zorder_nonoverlap``: 35 of 39 evaluable mutants checked by
``compare_pptx_files``; 29 confirmed by both raters, 6 unresolved). The correction changes
only that, with the operator's own notion of "cannot paint the same pixel"
(``harness.q2_mutation.operators.pptx.zorder_pair_is_inert``):

* a shape's box is its own ``a:off``/``a:ext`` frame, widened to the bounding box of its
  rotation; a shape without its own offset and extent (a placeholder that inherits its
  layout's position) has no known box;
* two boxes are apart only if neither comes within 2 mm (``CLEARANCE_EMU`` = 72,000 EMU) of
  the other; an unknown box is never apart from anything.

For each slide pair (by index), the agent's deck (``file1``) is aligned to the reference
(``file2``): for each reference shape in order, the first not-yet-placed agent shape with
the same signature (element kind, shape type, frame, every text run) is taken, and it may
move ahead of the agent shapes it skips only if it is apart from each of them. So the
alignment changes the relative order of a pair only when the pair is apart, which is exactly
the swaps the operator calls equivalent; overlapping shapes keep their order and a visible
stacking change still fails. If any reference shape finds no such agent shape, that slide is
left as it is. If no slide changed, the original comparator runs on the original file, so
its verdict is the original verdict; otherwise the realigned deck is written to a temporary
file and the original comparator runs on it with the same options.

Nothing else is corrected: ``compare_pptx_files_tolerant`` calls the slides module's own
function and keeps its z-order false negative (task ``a434992a``, flagged), and confirmed
false positives stay (section 8). The original comparator is imported from the pinned OSWorld
``b138d348`` tree at call time, so this module needs python-pptx only.
"""

from __future__ import annotations

import math
import os
import tempfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from typing import Any

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
CLEARANCE_EMU = 72_000  # 2 mm, as harness.q2_mutation.operators.pptx.CLEARANCE_EMU

Box = tuple[int, int, int, int]


@dataclass(frozen=True)
class ShapeKey:
    """What identifies a shape across the two decks, and its box (EMU) if known."""

    signature: tuple[Any, ...]
    box: Box | None


_XFRM_PATHS = {
    "sp": f"{P_NS}spPr/{A_NS}xfrm",
    "cxnSp": f"{P_NS}spPr/{A_NS}xfrm",
    "pic": f"{P_NS}spPr/{A_NS}xfrm",
    "grpSp": f"{P_NS}grpSpPr/{A_NS}xfrm",
    "graphicFrame": f"{P_NS}xfrm",
}


def _xfrm(element: Any) -> Any:
    """The shape's own transform (``spPr``/``grpSpPr`` ``a:xfrm``, or ``p:xfrm``), or None."""
    path = _XFRM_PATHS.get(element.tag.rsplit("}", 1)[-1])
    return None if path is None else element.find(path)


def shape_box(element: Any) -> Box | None:
    """The frame from the element's own offset and extent, widened by its rotation."""
    xfrm = _xfrm(element)
    if xfrm is None:
        return None
    off = xfrm.find(f"{A_NS}off")
    ext = xfrm.find(f"{A_NS}ext")
    if off is None or ext is None:
        return None
    try:
        x, y = int(off.get("x")), int(off.get("y"))
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        rot = int(xfrm.get("rot") or 0)
    except (TypeError, ValueError):
        return None
    if rot:
        angle = math.radians(rot / 60_000)
        w = abs(cx * math.cos(angle)) + abs(cy * math.sin(angle))
        h = abs(cx * math.sin(angle)) + abs(cy * math.cos(angle))
        mx, my = x + cx / 2, y + cy / 2
        return (math.floor(mx - w / 2), math.floor(my - h / 2),
                math.ceil(mx + w / 2), math.ceil(my + h / 2))  # fmt: skip
    return x, y, x + cx, y + cy


def apart(a: ShapeKey, b: ShapeKey) -> bool:
    """Neither box comes within ``CLEARANCE_EMU`` of the other (unknown boxes never are)."""
    if a.box is None or b.box is None:
        return False
    l1, t1, r1, b1 = a.box
    l2, t2, r2, b2 = b.box
    c = CLEARANCE_EMU
    return r1 + c <= l2 or r2 + c <= l1 or b1 + c <= t2 or b2 + c <= t1


def shape_key(shape: Any) -> ShapeKey:
    element = shape._element
    tag = element.tag.rsplit("}", 1)[-1]
    texts = tuple(node.text or "" for node in element.iter(f"{A_NS}t"))
    try:
        kind = int(shape.shape_type) if shape.shape_type is not None else None
    except Exception:  # noqa: BLE001 - unrecognized graphic frames raise NotImplementedError
        kind = None
    box = shape_box(element)
    return ShapeKey(signature=(tag, kind, box, texts), box=box)


def align_order(source: Sequence[ShapeKey], target: Sequence[ShapeKey]) -> list[int] | None:
    """The permutation (indices into ``source``) that puts ``source`` in ``target``'s order,
    moving a shape only past shapes it is apart from; None when there is none."""
    if len(source) != len(target):
        return None
    remaining = list(range(len(source)))
    order: list[int] = []
    for want in target:
        pick = None
        for position, index in enumerate(remaining):
            if source[index].signature != want.signature:
                continue
            skipped = remaining[:position]
            if all(apart(source[index], source[other]) for other in skipped):
                pick = position
            break  # the first signature match only: identical shapes keep their order
        if pick is None:
            return None
        order.append(remaining.pop(pick))
    return order


def _reorder_slide(slide: Any, order: Sequence[int]) -> None:
    """Put the slide's top-level shape elements into ``order`` (other children stay put)."""
    tree = slide.shapes._spTree
    elements = [shape._element for shape in slide.shapes]
    slots = [i for i, child in enumerate(tree) if child in elements]
    for element in elements:
        tree.remove(element)
    for slot, index in sorted(zip(slots, order, strict=True)):
        tree.insert(slot, elements[index])


def aligned_deck(file1_path: str, file2_path: str) -> tuple[Any, list[dict[str, Any]]] | None:
    """The agent deck with its apart shapes put in the reference's order, and a per-slide
    report; None when no slide changes (or the decks cannot be read)."""
    from pptx import Presentation

    try:
        prs1 = Presentation(file1_path)
        prs2 = Presentation(file2_path)
    except Exception:  # noqa: BLE001 - the original comparator raises or fails on its own
        return None
    report: list[dict[str, Any]] = []
    changed = False
    for index, (slide1, slide2) in enumerate(zip(prs1.slides, prs2.slides, strict=False), 1):
        source = [shape_key(shape) for shape in slide1.shapes]
        target = [shape_key(shape) for shape in slide2.shapes]
        order = align_order(source, target)
        if order is None:
            report.append({"slide": index, "aligned": False})
            continue
        moved = order != list(range(len(source)))
        report.append({"slide": index, "aligned": True, "reordered": moved})
        if moved:
            _reorder_slide(slide1, order)
            changed = True
    if not changed:
        return None
    return prs1, report


def original_comparator() -> Callable[..., Any]:
    from desktop_env.evaluators.metrics.slides import compare_pptx_files

    return compare_pptx_files


def compare_pptx_files_zinv(
    file1_path: str,
    file2_path: str,
    _original: Callable[..., Any] | None = None,
    **options: Any,
) -> Any:
    """``compare_pptx_files`` after aligning the agent deck's apart shapes."""
    original = _original or original_comparator()
    if file1_path is None or file2_path is None:
        return original(file1_path, file2_path, **options)
    try:
        aligned = aligned_deck(file1_path, file2_path)
    except Exception:  # noqa: BLE001 - alignment never decides a verdict on its own
        aligned = None
    if aligned is None:
        return original(file1_path, file2_path, **options)
    prs1, _report = aligned
    handle, path = tempfile.mkstemp(prefix="zinv-", suffix=".pptx")
    os.close(handle)
    try:
        prs1.save(path)
        return original(path, file2_path, **options)
    finally:
        os.unlink(path)


def install(metrics_module: Any) -> Callable[..., Any]:
    """Replace ``compare_pptx_files`` in OSWorld's ``metrics`` package namespace (what
    ``DesktopEnv._set_evaluator_info`` resolves by name) and return the original. The tolerant
    variant calls the slides module's own function and is therefore not corrected."""
    original = metrics_module.compare_pptx_files
    metrics_module.compare_pptx_files = compare_pptx_files_zinv
    return original
